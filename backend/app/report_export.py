"""Excel (and CSV) reports built from the same authorized data as the panel — portfolio or a single work."""
import csv
import io
from datetime import datetime, timezone
from decimal import Decimal

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .alerts import detect
from .db import get_db
from .intelligence import project_intelligence
from .materials import material_rows
from .measurements import access
from .models import Expense, LaborContract, Material, Measurement, Membership, Payment, Project, ServiceItem
from .security import require_identity

router = APIRouter(prefix="/api/v1", tags=["reports"])
READ_ROLES = ["OWNER", "ADMIN", "ENGINEER", "INSPECTOR", "FIELD_OPERATOR", "VIEWER"]
CATEGORY = {"MATERIAL": "Material", "LABOR": "Mão de obra", "SERVICE": "Serviço", "FUEL": "Combustível", "EQUIPMENT": "Equipamento", "OTHER": "Outros"}
STATUS = {"PLANNED": "Planejada", "IN_PROGRESS": "Em execução", "PAUSED": "Paralisada", "DONE": "Concluída", "COST_POOL": "A ratear"}
MONEY, QTY, PCT, DATE = 'R$ #,##0.00', '#,##0.000', '0.0"%"', "dd/mm/yyyy"
HEADER = PatternFill("solid", fgColor="0E7C86")


def _sheet(book: Workbook, title: str, columns: list[tuple[str, str | None, int]], rows: list[list]):
    sheet = book.create_sheet(title)
    sheet.append([name for name, _, _ in columns])
    for cell in sheet[1]:
        cell.font, cell.fill, cell.alignment = Font(bold=True, color="FFFFFF"), HEADER, Alignment(vertical="center", wrap_text=True)
    for row in rows:
        sheet.append(row)
    for index, (_, fmt, width) in enumerate(columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = width
        if fmt:
            for cell in sheet[letter][1:]:
                cell.number_format = fmt
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def _d(value):
    return None if value is None else float(value)


def build_workbook(db: Session, projects: list[Project], subject: str, scope: str) -> bytes:
    ids = [p.id for p in projects]
    names = {p.id: p.name for p in projects}
    book = Workbook()
    book.remove(book.active)
    paid = {eid: total for eid, total in db.execute(select(Payment.expense_id, func.sum(Payment.amount)).where(Payment.status == "PAID")
                                                     .group_by(Payment.expense_id)).all()}
    expenses = db.scalars(select(Expense).where(Expense.project_id.in_(ids)).order_by(Expense.expense_date)).all() if ids else []

    summary = []
    for p in projects:
        intel = project_intelligence(db, p) if p.status != "COST_POOL" else None
        s, c = (intel["schedule"], intel["cost"]) if intel else ({}, {})
        incurred = sum((e.amount for e in expenses if e.project_id == p.id), Decimal(0))
        paid_total = sum((paid.get(e.id, Decimal(0)) for e in expenses if e.project_id == p.id), Decimal(0))
        summary.append([p.name, STATUS.get(p.status, p.status), p.municipality, p.responsible, _d(p.budget), s.get("current_progress"),
                        _d(incurred), _d(paid_total), _d(incurred - paid_total), s.get("finish_p50"), s.get("finish_p10"), s.get("finish_p90"),
                        None if s.get("probability_on_time") is None else round(s["probability_on_time"] * 100, 1),
                        c.get("eac_p50"), c.get("cpi"), "; ".join(s.get("missing") or c.get("missing") or []) or None])
    _sheet(book, "Resumo", [("Obra", None, 34), ("Situação", None, 14), ("Município", None, 14), ("Responsável", None, 20),
                            ("Orçamento", MONEY, 16), ("Avanço físico (%)", PCT, 12), ("Custo incorrido", MONEY, 16), ("Pago", MONEY, 16),
                            ("Em aberto", MONEY, 16), ("Término provável (P50)", None, 14), ("Término P10", None, 12), ("Término P90", None, 12),
                            ("Prob. no prazo (%)", PCT, 12), ("Custo final estimado", MONEY, 18), ("CPI", "0.000", 8),
                            ("O que falta para prever", None, 50)], summary)

    _sheet(book, "Despesas", [("Obra", None, 30), ("Data", DATE, 12), ("Fornecedor", None, 26), ("Descrição", None, 40), ("Categoria", None, 14),
                              ("Valor", MONEY, 14), ("Pago", MONEY, 14), ("Em aberto", MONEY, 14), ("Observação", None, 40), ("Origem", None, 26)],
           [[names[e.project_id], e.expense_date, e.supplier, e.description, CATEGORY.get(e.category, e.category), _d(e.amount),
             _d(paid.get(e.id, 0)), _d(e.amount - paid.get(e.id, Decimal(0))), e.notes, e.source_ref] for e in expenses])

    payments = db.execute(select(Payment, Expense).join(Expense, Payment.expense_id == Expense.id).where(Expense.project_id.in_(ids))
                          .order_by(Payment.scheduled_date)).all() if ids else []
    _sheet(book, "Pagamentos", [("Obra", None, 30), ("Fornecedor", None, 26), ("Data", DATE, 12), ("Valor", MONEY, 14), ("Situação", None, 12),
                                ("Comprovante", None, 30)],
           [[names[e.project_id], e.supplier, pay.scheduled_date, _d(pay.amount), "Pago" if pay.status == "PAID" else "Programado", pay.reference]
            for pay, e in payments])

    contracts = db.scalars(select(LaborContract).where(LaborContract.project_id.in_(ids))).all() if ids else []
    _sheet(book, "Contratos", [("Obra", None, 30), ("Empreiteiro", None, 28), ("Contrato", MONEY, 16), ("Adiantado (informado)", MONEY, 18),
                               ("Saldo (informado)", MONEY, 16)],
           [[names[c.project_id], c.contractor, _d(c.total), _d(c.reported_advance), _d(c.reported_balance)] for c in contracts])

    measurements = db.execute(select(Measurement, ServiceItem).join(ServiceItem, Measurement.service_item_id == ServiceItem.id)
                              .where(ServiceItem.project_id.in_(ids)).order_by(Measurement.measured_on)).all() if ids else []
    _sheet(book, "Medições", [("Obra", None, 30), ("Serviço", None, 30), ("Data", DATE, 12), ("Quantidade", QTY, 12), ("Unidade", None, 8),
                              ("Situação", None, 12), ("Registrado por", None, 18), ("Decidido por", None, 18), ("Motivo", None, 30)],
           [[names[s.project_id], s.name, m.measured_on, _d(m.quantity), s.unit,
             {"PENDING": "Pendente", "APPROVED": "Aprovada", "REJECTED": "Recusada"}[m.status], m.recorded_by, m.decided_by, m.decision_reason]
            for m, s in measurements])

    materials = db.scalars(select(Material).where(Material.project_id.in_(ids))).all() if ids else []
    _sheet(book, "Materiais", [("Obra", None, 30), ("Material", None, 26), ("Unidade", None, 8), ("Saldo", QTY, 12), ("Custo médio", MONEY, 14),
                               ("Valor em estoque", MONEY, 16), ("Consumido", QTY, 12), ("Esperado pelo avanço", QTY, 14),
                               ("Cobertura (dias)", "0", 12), ("Sinais", None, 30)],
           [[names[r["project_id"]], r["name"], r["unit"], _d(r["balance"]), _d(r["avg_cost"]), _d(r["stock_value"]), _d(r["consumed"]),
             _d(r["expected_consumption"]), r["coverage_days"], ", ".join(r["flags"]) or None] for r in material_rows(db, list(materials))])

    alerts = detect(db, projects)
    _sheet(book, "Alertas", [("Prioridade", None, 10), ("Tipo", None, 22), ("Obra", None, 30), ("Alerta", None, 44), ("Detalhe", None, 70)],
           [[{"high": "Alta", "medium": "Média", "low": "Baixa"}[a["severity"]], a["type"], a["project"], a["title"], a["detail"]] for a in alerts])

    about = book.create_sheet("Sobre")
    for line in (["Aggilli Obras Intelligence — relatório " + scope],
                 ["Gerado em", datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")], ["Por", subject],
                 ["Previsões", "Monte Carlo de produtividade e Earned Schedule (término); Valor Agregado/PMI ou custo por avanço (custo final)."],
                 ["Atenção", "Previsões são estimativas com faixa de incerteza; pagamentos programados não são pagamentos confirmados; "
                             "saldos de contratos são os informados na fonte."]):
        about.append(line)
    about["A1"].font = Font(bold=True, size=14)
    about.column_dimensions["A"].width, about.column_dimensions["B"].width = 16, 110
    db.commit()  # forecast runs recorded while building the report
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _xlsx(content: bytes, filename: str) -> StreamingResponse:
    return StreamingResponse(iter([content]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.get("/reports/portfolio.xlsx")
def portfolio_xlsx(subject: str = Depends(require_identity), db: Session = Depends(get_db)):
    projects = list(db.scalars(select(Project).join(Membership, Membership.organization_id == Project.organization_id)
                               .where(Membership.subject_id == subject, Membership.role.in_(READ_ROLES)).order_by(Project.name)).all())
    return _xlsx(build_workbook(db, projects, subject, "da carteira"), "carteira-obras.xlsx")


@router.get("/projects/{project_id}/reports/obra.xlsx")
def project_xlsx(project_id: str, subject: str = Depends(require_identity), db: Session = Depends(get_db)):
    project = access(db, subject, project_id, READ_ROLES)
    return _xlsx(build_workbook(db, [project], subject, "da obra " + project.name), "obra.xlsx")


@router.get("/projects/{project_id}/reports/expenses.csv")
def expenses_csv(project_id: str, subject: str = Depends(require_identity), db: Session = Depends(get_db)):
    project = access(db, subject, project_id, READ_ROLES)
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["Data", "Fornecedor", "Descrição", "Categoria", "Valor", "Observação"])
    for e in db.scalars(select(Expense).where(Expense.project_id == project.id).order_by(Expense.expense_date)).all():
        writer.writerow([e.expense_date.strftime("%d/%m/%Y") if e.expense_date else "", e.supplier, e.description,
                         CATEGORY.get(e.category, e.category), f"{e.amount:.2f}".replace(".", ","), e.notes or ""])
    return StreamingResponse(iter(["﻿" + output.getvalue()]), media_type="text/csv; charset=utf-8",
                             headers={"Content-Disposition": 'attachment; filename="despesas.csv"'})
