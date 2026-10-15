"""Excel/CSV reports: all sheets present, figures match the database, scope respects authorization."""
import io
import json

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.db import Base, get_db
from app.main import app
from app.models import Project
from app.security import require_identity
from scripts import import_tuntum


def test_portfolio_and_project_reports():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        import_tuntum.run(db, "Grupo MG", "owner", "import:test", json.loads(import_tuntum.DATA.read_text(encoding="utf-8")))
        db.commit()
        pid = db.scalar(select(Project.id).where(Project.name == "MG Tabocal"))

    def get_test_db():
        with Session(engine) as session:
            yield session

    app.dependency_overrides[get_db] = get_test_db
    who = {"sub": "owner"}
    app.dependency_overrides[require_identity] = lambda: who["sub"]
    try:
        client = TestClient(app)
        response = client.get("/api/v1/reports/portfolio.xlsx")
        assert response.status_code == 200 and response.headers["content-type"].startswith("application/vnd.openxml")
        book = load_workbook(io.BytesIO(response.content))
        assert book.sheetnames == ["Resumo", "Despesas", "Pagamentos", "Contratos", "Medições", "Materiais", "Alertas", "Sobre"]
        payments = [row for row in book["Pagamentos"].iter_rows(min_row=2, values_only=True)]
        assert round(sum(row[3] for row in payments), 2) == 31570.87
        assert book["Resumo"].max_row - 1 == 15  # 14 works + cost pool
        obra = load_workbook(io.BytesIO(client.get(f"/api/v1/projects/{pid}/reports/obra.xlsx").content))
        assert {row[0] for row in obra["Despesas"].iter_rows(min_row=2, values_only=True)} == {"MG Tabocal"}
        csv_text = client.get(f"/api/v1/projects/{pid}/reports/expenses.csv").text
        assert csv_text.startswith("﻿Data;Fornecedor") and "2000,00" in csv_text
        who["sub"] = "intruso"
        assert client.get(f"/api/v1/projects/{pid}/reports/obra.xlsx").status_code == 404
        empty = load_workbook(io.BytesIO(client.get("/api/v1/reports/portfolio.xlsx").content))
        assert empty["Resumo"].max_row == 1
    finally:
        app.dependency_overrides.clear()
        engine.dispose()
